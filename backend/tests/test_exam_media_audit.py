"""Synthetic Word image fixtures; never the learner's copyrighted exam files."""
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from docx import Document
from PIL import Image

from app.exam_media_audit import audit_bytes


def synthetic_docx(tmp_path, *, picture=True):
    document = Document()
    document.add_paragraph("Synthetic exam figure fixture, no copyrighted text")
    if picture:
        image = tmp_path / "test.png"
        Image.new("RGB", (4, 4), "white").save(image)
        document.add_picture(str(image))
    target = tmp_path / "synthetic.docx"
    document.save(target)
    return target.read_bytes()


def repackage(raw, *, removed=None, extras=None):
    removed = set(removed or [])
    extras = extras or {}
    output = BytesIO()
    with ZipFile(BytesIO(raw)) as src, ZipFile(output, "w", ZIP_DEFLATED) as dst:
        for member in src.infolist():
            if member.filename not in removed:
                dst.writestr(member, src.read(member.filename))
        for filename, data in extras.items():
            dst.writestr(filename, data)
    return output.getvalue()


def test_image_is_hashed_and_mapped_to_word_paragraph_without_content(tmp_path):
    result = audit_bytes(synthetic_docx(tmp_path))
    assert result["media_objects"] == result["referenced_media_objects"] == 1
    assert result["unreferenced_media_objects"] == []
    assert result["media"][0]["paragraph_ordinals"]
    assert len(result["media"][0]["sha256"]) == 64
    assert "Synthetic exam figure" not in str(result)
    assert result["images_scientifically_verified"] is False
    assert result["question_image_anchors_verified"] is False


def test_no_media_never_asserts_visual_scientific_review(tmp_path):
    result = audit_bytes(synthetic_docx(tmp_path, picture=False))
    assert result["media_objects"] == 0
    assert result["images_scientifically_verified"] is False


def test_unanchored_word_media_remains_reported_as_risk(tmp_path):
    blob = repackage(
        synthetic_docx(tmp_path), extras={"word/media/hdphoto1.wdp": b"test"}
    )
    report = audit_bytes(blob)
    assert report["media_objects"] == 2
    assert report["referenced_media_objects"] == 1
    assert report["unreferenced_media_objects"] == ["word/media/hdphoto1.wdp"]


def test_missing_embedded_image_blob_is_rejected(tmp_path):
    raw = synthetic_docx(tmp_path)
    with ZipFile(BytesIO(raw)) as archive:
        media_name = next(
            x for x in archive.namelist() if x.startswith("word/media/")
        )
    with pytest.raises(ValueError, match="docx_image_relationship_blob_missing"):
        audit_bytes(repackage(raw, removed={media_name}))


def test_non_docx_is_rejected_without_guessing():
    with pytest.raises(ValueError, match="invalid_exam_docx_media_archive"):
        audit_bytes(b"not a Word archive")
