"""Synthetic format/regression tests, never real scientific approval."""

import io
import hashlib
import pytest
from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
from psycopg.errors import RaiseException
from app.evidence_supplements import extract, math_text, annotate, assess
from test_ocr_only import import_book
from test_candidate_review import packet  # noqa: F401


def test_math_fraction_and_subscript_no_execution():
    node = etree.fromstring(
        b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:f><m:num><m:sSub><m:e><m:r><m:t>I</m:t></m:r></m:e><m:sub><m:r><m:t>0</m:t></m:r></m:sub></m:sSub></m:num><m:den><m:r><m:t>I</m:t></m:r></m:den></m:f></m:oMath>'
    )
    assert math_text(node) == "(I_{0})/(I)"


def test_extract_original_omml_and_nested_table_positions():
    doc = Document()
    doc.add_paragraph("SYNTHETIC context")
    p = doc.add_paragraph()
    m = OxmlElement("m:oMath")
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = "x=1"
    r.append(t)
    m.append(r)
    p._p.append(m)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0]._p.append(etree.fromstring(etree.tostring(m)))
    buf = io.BytesIO()
    doc.save(buf)
    rows = list(extract(buf.getvalue()))
    assert [x["position"] for x in rows] == ["body/2", "body/3/row1/cell1/1"]
    assert all(x["search_text"] == "x=1" and x["kind"] == "omml" for x in rows)
    assert hashlib.sha256(rows[0]["payload"]).digest()


def test_math_delimiters_and_function_argument_preserved():
    xml = '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:func><m:fName><m:r><m:t>log</m:t></m:r></m:fName><m:e><m:d><m:dPr><m:begChr m:val="["/><m:endChr m:val="]"/></m:dPr><m:e><m:r><m:t>S</m:t></m:r></m:e></m:d></m:e></m:func></m:oMath>'
    assert math_text(etree.fromstring(xml.encode())) == "log([S])"


def test_real_db_supplement_identity_search_and_no_overwrite(db, tmp_path):
    from test_ocr_only import material
    from app.sources import import_source
    from app.evidence_supplements import append_supplement

    path, meta = material(tmp_path, "biochemistry_exams")
    doc = Document(path)
    m = etree.fromstring(
        b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>syntheticFormula=1</m:t></m:r></m:oMath>'
    )
    doc.paragraphs[2]._p.append(m)
    doc.save(path)
    did = import_source(db, path, meta)
    source_sha = hashlib.sha256(path.read_bytes()).hexdigest()
    item = next(x for x in extract(path.read_bytes()) if x["kind"] == "omml")
    uid = db.execute(
        "select id from source_units where document_id=%s and section like 'body/3%%' order by id limit 1",
        (did,),
    ).fetchone()["id"]
    before = db.execute("select count(*) AS n from source_units").fetchone()["n"]
    key = append_supplement(db, did, source_sha, item, uid)
    assert append_supplement(db, did, source_sha, item, uid) == key
    assert (
        db.execute(
            "select search_text from formula_search_renderings where supplement_key=%s",
            (key,),
        ).fetchone()["search_text"]
        == "syntheticFormula=1"
    )
    assert (
        db.execute("select count(*) AS n from source_units").fetchone()["n"] == before
    )
    with pytest.raises(ValueError):
        append_supplement(db, did, source_sha, dict(item, payload=b"fabricated"), uid)
    with pytest.raises(RaiseException), db.transaction():
        db.execute("delete from source_supplements")


def test_scoped_science_pass_cannot_publish_and_is_immutable(db, packet):  # noqa: F811
    from app.candidate_review import create_packet

    qid = create_packet(db, packet)
    analysis = {
        k: "Synthetic evidence-bound analysis, not an independent automated round"
        for k in (
            "stem",
            "answer",
            "rubric",
            "science",
            "style",
            "media",
            "duplicates",
            "uncertainties",
        )
    }
    assess(db, "synthetic-science", packet["candidate_key"], "pass", analysis)
    assess(db, "synthetic-science", packet["candidate_key"], "pass", analysis)
    row = db.execute("SELECT * FROM candidate_science_assessments").fetchone()
    assert row["formal_gate_passed"] is False
    assert (
        db.execute(
            "SELECT status FROM reviewed_questions WHERE id=%s", (qid,)
        ).fetchone()["status"]
        == "draft"
    )
    assert (
        db.execute("SELECT count(*) AS n FROM auto_quality_events").fetchone()["n"] == 0
    )
    with pytest.raises(ValueError):
        assess(db, "synthetic-science", packet["candidate_key"], "reject", analysis)
    with pytest.raises(RaiseException), db.transaction():
        db.execute("UPDATE candidate_science_assessments SET conclusion='reject'")


def test_append_annotation_immutable_idempotent_and_not_source_approval(db, tmp_path):
    _, _, doc, _ = import_book(db, tmp_path)
    uid = db.execute(
        "select id from source_units where document_id=%s order by id", (doc,)
    ).fetchone()["id"]
    annotate(
        db,
        "synthetic-test",
        uid,
        "type_analysis",
        {"inferred": "term", "official": False},
    )
    annotate(
        db,
        "synthetic-test",
        uid,
        "type_analysis",
        {"inferred": "term", "official": False},
    )
    assert (
        db.execute("select count(*) AS n from source_text_annotations").fetchone()["n"]
        == 1
    )
    with pytest.raises(ValueError):
        annotate(db, "synthetic-test", uid, "type_analysis", {"inferred": "essay"})
    with pytest.raises(RaiseException), db.transaction():
        db.execute("delete from source_text_annotations")
    assert (
        db.execute("select count(*) AS n from source_review_events").fetchone()["n"]
        == 0
    )


def test_formula_search_requires_query_and_bounded_limit(db):
    from app.evidence_supplements import search_formulas

    for query, limit in [("", 20), ("x", 0), ("x", 101)]:
        with pytest.raises(ValueError):
            search_formulas(db, query, limit)
    assert search_formulas(db, "nonexistent") == []
