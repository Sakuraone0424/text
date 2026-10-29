"""Literal chapter↔genuine-exam linkage is only a review queue (synthetic fixtures)."""
import json
from hashlib import sha256

import pytest
from docx import Document

from app.offline_topic_linkage import report, science_terms
from app.source_policy import COLLECTION
from test_collection_seed import synthetic_collection


def test_terms_are_real_headings_not_invented_definitions():
    assert "合成测试" in science_terms("第一章 合成测试")
    assert "氧化磷酸化" in science_terms("第十九章 氧化磷酸化")
    assert science_terms("细胞") == []
    assert science_terms("生物化学") == []


def test_synthetic_literal_question_chapter_match_is_candidate_only(tmp_path):
    root, lock = synthetic_collection(tmp_path)
    item = next(x for x in COLLECTION["files"] if x["key"] == "cell_biology_exams")
    path = root / item["relative_path"]
    doc = Document(path)
    doc.paragraphs[2].text = "1. 合成测试是否能映射到来源章节？"
    doc.save(path)
    manifest = json.loads(lock.read_text("utf-8"))
    raw = path.read_bytes()
    manifest["files"][item["key"]].update(
        sha256=sha256(raw).hexdigest(), bytes=len(raw)
    )
    lock.write_text(json.dumps(manifest), encoding="utf-8")
    data = report(root, lock_path=lock)
    row = next(
        b for b in data["buckets"]
        if b["subject"] == "cell_biology"
        and b["question_type"] == "single_choice"
    )
    assert row["lexically_linked_question_starts"] >= 1
    assert row["candidate_chapter_links"] >= 1
    assert row["scientifically_supported_matches"] == 0
    assert data["chapter_linkage_fully_verified"] is False
    assert data["independent_ai_answer_review_performed"] is False
    assert data["question_answers_or_rubrics_generated"] == 0
    assert "合成测试是否能映射" not in json.dumps(data, ensure_ascii=False)


def test_unpinned_exam_revision_fails_before_topic_alignment(tmp_path):
    root, lock = synthetic_collection(tmp_path)
    item = next(x for x in COLLECTION["files"] if x["key"] == "cell_biology_exams")
    with (root / item["relative_path"]).open("ab") as stream:
        stream.write(b"changed exam bytes")
    with pytest.raises(ValueError, match="SHA-256"):
        report(root, lock_path=lock)


def test_same_table_uncertainty_blocks_offline_topic_and_live_style_gate(tmp_path):
    from app.offline_topic_linkage import REFUSE
    from app.transcript_analysis import STYLE_SOURCE_BLOCKERS
    from app.exam_style_readiness import BOUNDARY_ISSUES

    assert REFUSE == STYLE_SOURCE_BLOCKERS | BOUNDARY_ISSUES
    root, lock = synthetic_collection(tmp_path)
    item = next(x for x in COLLECTION["files"] if x["key"] == "cell_biology_exams")
    path = root / item["relative_path"]
    doc = Document()
    doc.add_paragraph("2025年中国科学技术大学细胞生物学真题")
    doc.add_paragraph("单项选择题")
    doc.add_paragraph("1. 合成测试需依据下一表结果回答")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "未校验的合成表格"
    doc.add_paragraph("2. 合成测试的另一个独立问题")
    # Close the section: the synthetic fixture's table is truly tied only
    # to question one, not to the following question.
    doc.add_paragraph("2026年中国科学技术大学细胞生物学真题")
    doc.save(path)
    manifest = json.loads(lock.read_text(encoding="utf-8"))
    raw = path.read_bytes()
    manifest["files"][item["key"]].update(
        sha256=sha256(raw).hexdigest(), bytes=len(raw)
    )
    lock.write_text(json.dumps(manifest), encoding="utf-8")
    data = report(root, lock_path=lock)
    row = next(
        b for b in data["buckets"]
        if b["subject"] == "cell_biology"
        and b["question_type"] == "single_choice"
    )
    assert row["observed_candidate_starts"] == 2
    assert row["blocked_before_topic_linkage"] == 1
    assert row["parser_safe_style_starts"] == 1
    assert row["lexically_linked_question_starts"] == 1
    assert row["scientifically_supported_matches"] == 0
    assert data["question_answers_or_rubrics_generated"] == 0
