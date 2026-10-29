"""Six-source report tests in rollback-only PostgreSQL fixture schemas.

Synthetic samples are NOT the learner's real six source files.
"""

from app import sources
from app.source_readiness import audit
from app.source_policy import COLLECTION
from test_ocr_only import material
from test_api_v2 import LEARNER


def test_six_unimported_sources_are_explicitly_unknown(db, client):
    report = audit(db)
    assert report["expected_files"] == 6
    assert report["imported_files"] == 0
    assert report["integrity_passed_files"] == 0
    assert report["complete_for_generation"] is False
    assert all(x["status"] == "not_imported" for x in report["files"])
    assert all(x["source_sha256"] is None for x in report["files"])
    assert all(x["chapter_unmapped_ocr_pages"] == [] for x in report["files"])
    response = client.get("/v2/quality/six-source-readiness", headers=LEARNER)
    assert response.status_code == 200
    assert len(response.json()["files"]) == 6


def test_synthetic_textbook_source_report_has_real_hash_and_page_checks(db, tmp_path):
    path, metadata = material(tmp_path, "cell_biology")
    document_id = sources.import_source(db, path, metadata)
    report = audit(db)
    book = next(x for x in report["files"] if x["key"] == "cell_biology")
    assert book["document_id"] == document_id
    assert book["status"] == "integrity_passed"
    assert book["expected_ocr_pages"] == 385
    assert book["ocr_page_markers"] == 385
    assert book["ocr_unit_count"] > 0
    assert book["missing_ocr_page_markers"] == []
    assert book["blank_ocr_pages"] == []
    assert len(book["source_sha256"]) == 64
    assert book["mapped_chapter_count"] >= 1
    assert book["mapped_chapter_ocr_pages"] == 385
    assert book["chapter_unmapped_ocr_pages"] == []
    assert book["chapter_unmapped_ocr_page_count"] == 0
    assert book["chapter_mapping_scientifically_verified"] is False
    assert book["printed_original_verified"] is False
    assert report["complete_for_generation"] is False
    assert all(x["status"] == "not_imported" for x in report["files"] if x["key"] != "cell_biology")


def test_risky_ocr_text_is_not_declared_clean(db, tmp_path):
    path, metadata = material(tmp_path, "biochemistry_notes", content="如图所示 ATP → ADP")
    sources.import_source(db, path, metadata)
    report = audit(db)
    notes = next(x for x in report["files"] if x["key"] == "biochemistry_notes")
    assert notes["status"] == "integrity_passed_with_ocr_risks"
    assert notes["risk_flag_counts"]["image_dependent"] >= 1
    assert notes["expected_ocr_pages"] == 289
    assert notes["printed_original_verified"] is False


def test_synthetic_exam_transcript_reports_transcriptions_not_official_original(db, tmp_path):
    path, metadata = material(tmp_path, "cell_biology_exams")
    sources.import_source(db, path, metadata)
    doc = next(x for x in audit(db)["files"] if x["key"] == "cell_biology_exams")
    assert doc["document_id"] is not None
    assert doc["exam_candidate_count"] >= 1
    assert doc["expected_ocr_pages"] is None
    assert doc["past_exam_source_status"] == "user_designated_genuine_past_exam"
    assert doc["status"].startswith("integrity_passed")


def test_docx_embedded_image_is_counted_but_not_approved_as_transcribed(db, tmp_path):
    from docx import Document
    from PIL import Image
    png = tmp_path / "synthetic-diagram.png"
    Image.new("RGB", (16, 16), color="white").save(png)
    path, metadata = material(tmp_path, "biochemistry_exams")
    document = Document(path)
    document.add_paragraph("A synthetic exam question with a figure")
    document.add_picture(str(png))
    document.save(path)
    sources.import_source(db, path, metadata)
    row = next(x for x in audit(db)["files"] if x["key"] == "biochemistry_exams")
    assert row["embedded_media_objects"] >= 1
    assert row["embedded_media_referenced_objects"] >= 1
    assert row["embedded_media_unanchored_objects"] == []
    assert row["embedded_media_formats"]["png"] >= 1
    assert row["embedded_media_review_status"] == "embedded_media_not_independently_transcribed"
    assert row["past_exam_source_status"] == "user_designated_genuine_past_exam"
