"""Only synthetic text in rollback-only test schemas; never a real-file verification claim."""

import json
from copy import deepcopy
import pytest
from docx import Document
from psycopg.errors import RaiseException
from app import (
    sources,
    ocr_sources as ocr,
    ocr_reports as reports,
    exams,
    source_intake,
)
from app.source_policy import COLLECTION, FILES
from app.source_files import sha

NOTE = "Synthetic reviewer action only inside an isolated rollback-only automated test"


def material(tmp_path, key="biochemistry_lower", content=None, *, synthetic=False):
    item = next(i for i in FILES.values() if i["key"] == key)
    path = tmp_path / item["filename"]
    metadata = dict(
        file_id=item["file_id"],
        folder_id=COLLECTION["folder_id"],
        collection_version=COLLECTION["version"],
        source_mode="ocr_only",
        title="SYNTHETIC TEST FIXTURE",
        authors="Test author",
        edition="Test edition",
        subject=item["subject"],
        material_status="synthetic_test" if synthetic else "test_fixture",
    )
    if item["kind"] == "textbook":
        metadata.update(
            expected_pages=item["expected_pages"],
            volume_offsets=[0],
            volume_pages=[item["expected_pages"]],
        )
        label = (item.get("page_label") + " ") if item.get("page_label") else ""
        path.write_text(
            "\n".join(
                f"===== {label}page-{n:03} =====\n第一章 合成测试\n"
                + (content or f"Synthetic evidence paragraph {n}")
                + "\n"
                for n in range(1, item["expected_pages"] + 1)
            ),
            encoding="utf-8",
        )
    else:
        doc = Document()
        for text in content or [
            "2023年中国科学技术大学真题",
            "单项选择题",
            "1. Synthetic 2023 question",
            "参考答案：not official",
            "2024年中国科学技术大学真题",
            "简答题",
            "1. Synthetic 2024 question",
            "2025年中国科学技术大学真题",
            "名词解释",
            "1. Synthetic 2025 question",
            "2026年中国科学技术大学真题",
            "论述题",
            "1. Synthetic 2026 question",
        ]:
            doc.add_paragraph(text)
        doc.add_table(rows=1, cols=1).cell(
            0, 0
        ).text = "Uncertain transcript continuation in a table"
        doc.save(path)
    return path, metadata


def import_book(db, tmp_path, content=None):
    path, metadata = material(tmp_path, content=content)
    doc = sources.import_source(db, path, metadata)
    row = db.execute(
        "SELECT * FROM source_units WHERE document_id=%s ORDER BY ordinal LIMIT 1",
        (doc,),
    ).fetchone()
    return path, metadata, doc, row


def approve(db, unit, kind):
    return ocr.review_unit(db, unit, kind, "approved", "test", NOTE)


@pytest.mark.parametrize(
    "key",
    ["biochemistry_upper", "biochemistry_lower", "biochemistry_notes", "cell_biology"],
)
def test_corrected_textbook_full_marker_formats_and_no_pdf_claim(tmp_path, key):
    path, metadata = material(tmp_path, key)
    result = sources.prepare_source(path, metadata)
    expected = FILES[metadata["file_id"]]["expected_pages"]
    assert len(result["pages"]) == expected
    assert [p["ocr_global_page"] for p in result["pages"]] == list(
        range(1, expected + 1)
    )
    assert all(
        p["pdf_page"] is None and p["printed_page"] is None for p in result["pages"]
    )
    assert all(r["pdf_page"] is None for r in result["rows"])
    assert all(p["marker_text"].startswith("===== ") for p in result["pages"])
    if key == "cell_biology":
        assert (
            result["pages"][192]["ocr_page"] == 193
            and result["pages"][192]["volume"] == 1
        )


def test_exact_current_six_sources_total_1915_and_no_old_ids():
    assert len(FILES) == 6
    assert sum(f.get("expected_pages", 0) for f in FILES.values()) == 1915
    assert next(f for f in FILES.values() if f["key"] == "biochemistry_notes")[
        "file_id"
    ].endswith("-_")
    assert "1aipWuwHLXiwACC2DuyXbtp5iudS820_7" not in FILES


@pytest.mark.parametrize(
    "change",
    [
        lambda m: m.update(file_id="old-unsupported-id"),
        lambda m: m.update(folder_id="outside-folder"),
        lambda m: m.update(expected_pages=586),
        lambda m: m.update(volume_offsets=[0, 192], volume_pages=[192, 193]),
        lambda m: m.update(pdf_pages={"1": 1}),
        lambda m: m.update(printed_pages={"1": "i"}),
    ],
)
def test_ocr_source_identity_and_physical_claims_rejected(tmp_path, change):
    path, meta = material(tmp_path)
    change(meta)
    with pytest.raises(ValueError):
        sources.prepare_source(path, meta)


@pytest.mark.parametrize(
    "mutation",
    ["missing", "duplicate", "backward", "malformed", "preamble", "wrong_label"],
)
def test_labelled_actual_format_rejects_corruption(tmp_path, mutation):
    path, meta = material(tmp_path, "biochemistry_upper")
    text = path.read_text()
    if mutation == "missing":
        text = text.replace("===== 上册 page-002 =====\n", "")
    if mutation == "duplicate":
        text = text.replace("page-002", "page-001", 1)
    if mutation == "backward":
        text = text.replace("page-003", "page-002", 1)
    if mutation == "malformed":
        text = text.replace("page-002", "page-0x2", 1)
    if mutation == "preamble":
        text = "unlocated body\n" + text
    if mutation == "wrong_label":
        text = text.replace("上册", "下册", 1)
    path.write_text(text)
    with pytest.raises(ValueError):
        sources.prepare_source(path, meta)


def test_txt_is_original_bytes_marker_identity_and_separate_stages(db, tmp_path):
    path, meta, doc, row = import_book(db, tmp_path)
    version = db.execute(
        "SELECT * FROM source_versions WHERE document_id=%s", (doc,)
    ).fetchone()
    assert bytes(version["source_content"]) == path.read_bytes() and version[
        "source_sha256"
    ] == sha(path.read_bytes())
    assert sources.import_source(db, path, meta) == doc
    header = db.execute("SELECT * FROM source_documents WHERE id=%s", (doc,)).fetchone()
    assert header["original_content"] is None and header["original_sha256"] is None
    assert (
        db.execute(
            "SELECT marker_text FROM source_pages WHERE document_id=%s ORDER BY ocr_global_page LIMIT 1",
            (doc,),
        ).fetchone()["marker_text"]
        == "===== 下册 page-001 ====="
    )
    state = reports.document_report(db, doc)
    assert not state["issues"] and state["stages"]["database_imported"]
    assert (
        state["stages"]["text_reviewed_units"] == 0
        and state["stages"]["science_reviewed_units"] == 0
    )
    assert (
        not state["pdf_physical_pages_verified"] and not state["printed_pages_verified"]
    )
    assert sources.search(db, "Synthetic evidence paragraph 1")
    approve(db, row["id"], "text")
    assert reports.document_report(db, doc)["stages"]["science_reviewed_units"] == 0
    approve(db, row["id"], "science")
    state = reports.document_report(db, doc)
    assert (
        state["stages"]["text_reviewed_units"] == 1
        and state["stages"]["science_reviewed_units"] == 1
    )
    assert state["chapters"]["第一章 合成测试"]["text_review_coverage"] < 1


def evidence(meta, doc, row, checksum):
    return dict(
        source_unit_id=row["id"],
        file_id=meta["file_id"],
        source_sha256=checksum,
        ocr_page=row["ocr_global_page"],
        paragraph_id=row["id"],
        textbook=meta["title"],
        authors=meta["authors"],
        edition=meta["edition"],
        isbn=None,
        chapter=row["chapter"],
        locator=f"OCR:{row['ocr_global_page']}:UNIT:{row['id']}",
        excerpt="Synthetic evidence paragraph 1",
    )


def test_ocr_question_requires_text_science_and_exact_citation(
    db, tmp_path, question_factory
):
    path, meta, doc, row = import_book(db, tmp_path)
    question = question_factory().model_dump()
    question["evidence"] = [evidence(meta, doc, row, sha(path.read_bytes()))]
    assert "unverified_or_missing_textbook_unit" in sources.evidence_issues(
        db, question
    )
    with pytest.raises(ValueError, match="Text review"):
        approve(db, row["id"], "science")
    approve(db, row["id"], "text")
    assert "ocr_science_not_reviewed" in sources.evidence_issues(db, question)
    approve(db, row["id"], "science")
    assert not sources.evidence_issues(db, question)
    for field, value in [
        ("file_id", "outside"),
        ("source_sha256", "0" * 64),
        ("paragraph_id", 99999),
        ("ocr_page", 2),
        ("locator", "PDF:1"),
    ]:
        bad = deepcopy(question)
        bad["evidence"][0][field] = value
        assert "ocr_evidence_identity_mismatch" in sources.evidence_issues(db, bad)
    ocr.review_unit(db, row["id"], "science", "rejected", "test", NOTE)
    assert "ocr_science_not_reviewed" in sources.evidence_issues(db, question)


def test_changed_file_gets_new_parent_version_without_reviews_or_links(db, tmp_path):
    path, meta, doc, row = import_book(db, tmp_path)
    approve(db, row["id"], "text")
    approve(db, row["id"], "science")
    sources.annotate(db, row["id"], ["synthetic/old-only"])
    path.write_text(path.read_text().replace("paragraph 1", "revised paragraph 1", 1))
    new = sources.import_source(db, path, meta)
    assert new != doc
    assert (
        db.execute(
            "SELECT previous_document_id FROM source_versions WHERE document_id=%s",
            (new,),
        ).fetchone()["previous_document_id"]
        == doc
    )
    assert not db.execute(
        "SELECT 1 FROM source_units WHERE document_id=%s AND verified", (new,)
    ).fetchone()
    assert not db.execute(
        "SELECT 1 FROM knowledge_links k JOIN source_units u ON u.id=k.unit_id WHERE u.document_id=%s",
        (new,),
    ).fetchone()
    with pytest.raises(ValueError, match="superseded"):
        approve(db, row["id"], "text")
    for query, args in [
        ("UPDATE source_units SET text=%s WHERE id=%s", ("tampered", row["id"])),
        ("UPDATE source_documents SET title=%s WHERE id=%s", ("tampered", doc)),
        ("DELETE FROM source_pages WHERE document_id=%s", (doc,)),
        ("DELETE FROM source_review_events WHERE unit_id=%s", (row["id"],)),
    ]:
        with pytest.raises(RaiseException), db.transaction():
            db.execute(query, args)


@pytest.mark.parametrize(
    "text",
    [
        "See 见图1 [图片]",
        "Chemical 化学结构式 here",
        "复杂表格 见表1",
        "uncertain ??? text",
        "ATP → ADP",
    ],
)
def test_uncertainty_needs_explicit_resolution_not_guessed_publication(
    db, tmp_path, text
):
    _, _, _, row = import_book(db, tmp_path, content=text)
    approve(db, row["id"], "text")
    with pytest.raises(ValueError, match="Unresolved"):
        approve(db, row["id"], "science")
    with pytest.raises(ValueError, match="flags"):
        ocr.review_unit(db, row["id"], "uncertainty", "approved", "test", NOTE, {})
    ocr.review_unit(
        db,
        row["id"],
        "uncertainty",
        "approved",
        "test",
        NOTE,
        dict(
            resolved_flags=row["quality_flags"],
            resolution="text_sufficient",
            explanation=NOTE,
        ),
    )
    approve(db, row["id"], "science")


def test_corrupt_and_synthetic_material_cannot_be_human_approved(db, tmp_path):
    _, _, _, row = import_book(db, tmp_path, content="corrupt \ufffd body")
    with pytest.raises(ValueError, match="Corrupt"):
        approve(db, row["id"], "text")
    path, meta = material(tmp_path, "biochemistry_notes", synthetic=True)
    doc = sources.import_source(db, path, meta)
    unit = db.execute(
        "SELECT id FROM source_units WHERE document_id=%s LIMIT 1", (doc,)
    ).fetchone()["id"]
    with pytest.raises(ValueError, match="Synthetic"):
        approve(db, unit, "text")


def test_genuine_past_exam_documents_count_recorded_questions_without_fake_answers(db, tmp_path):
    path, meta = material(tmp_path, "biochemistry_exams")
    doc = sources.import_source(db, path, meta)
    assert db.execute("SELECT count(*) AS n FROM exam_records").fetchone()["n"] == 0
    years = reports.annual_report(db)
    for year in range(2023, 2027):
        state = next(
            r for r in years if r["subject"] == "biochemistry" and r["year"] == year
        )
        assert state["candidate_questions"] == 1 and state["ocr_transcription_present"]
        assert state["frequency_eligible_from_transcription"]
        assert state["source_level"] == "user_provided_genuine_past_exam"
        assert state["textbook_answer_verification"] == "not_checked_in_year_coverage_report"
        assert state["priority_2023_2026"]
    for row in db.execute(
        "SELECT * FROM source_units WHERE document_id=%s", (doc,)
    ).fetchall():
        assert row["pdf_page"] is None and row["ocr_global_page"] is None
    unit = db.execute(
        "SELECT * FROM source_units WHERE document_id=%s AND text LIKE %s",
        (doc, "%Synthetic 2026 question%"),
    ).fetchone()
    approve(db, unit["id"], "text")
    approve(db, unit["id"], "exam_style")
    record = dict(
        source_unit_id=unit["id"],
        original_question_id="synthetic-2026",
        subject="biochemistry",
        year=2026,
        question_type="essay",
        stem=unit["text"],
        category="synthesis",
    )
    exams.add(db, record)
    for update in (
        {"original_verified": True},
        {"answer_kind": "official", "answer": "invented"},
    ):
        with pytest.raises(ValueError, match="transcription alone"):
            exams.add(db, {**record, **update})
    assert not any(
        r["trusted"]
        for r in exams.trusted_records(
            db, db.execute("SELECT clock_timestamp() AS now").fetchone()["now"]
        )
    )


def test_docx_unknown_year_and_range_title_do_not_fabricate_2026_coverage(db, tmp_path):
    path, meta = material(
        tmp_path,
        "cell_biology_exams",
        content=[
            "1998-2026年中国科学技术大学真题汇编",
            "单项选择",
            "1. Unassigned synthetic question",
            "答案：personal reference",
        ],
    )
    sources.import_source(db, path, meta)
    years = [r for r in reports.annual_report(db) if r["subject"] == "cell_biology"]
    assert all(r["candidate_questions"] == 0 for r in years)
    assert all(r["unassigned_year_paragraphs"] > 0 for r in years)


def test_cli_ocr_only_preflight_import_without_any_pdf(db, tmp_path, monkeypatch):
    from contextlib import contextmanager

    path, meta = material(tmp_path)
    metadata = tmp_path / "metadata.json"
    metadata.write_text(json.dumps(meta))
    args = [str(path), "--metadata", str(metadata)]
    pre = tmp_path / "preflight.json"
    assert source_intake.main(["check", *args, "--report", str(pre)]) == 0
    assert json.loads(pre.read_text())["stages"]["database_imported"] is False

    @contextmanager
    def isolated(*args, **kwargs):
        with db.transaction():
            yield db

    monkeypatch.setattr(source_intake, "isolated_db", isolated)
    imported = tmp_path / "imported.json"
    assert (
        source_intake.main(
            [
                "import",
                *args,
                "--schema",
                "source_staging_test",
                "--target",
                "development",
                "--report",
                str(imported),
            ]
        )
        == 0
    )
    state = json.loads(imported.read_text())
    assert state["database_audit"]["stages"]["text_reviewed_units"] == 0
    assert state["database_audit"]["stages"]["database_imported"]


def test_collection_status_distinguishes_user_location_from_actual_read(
    db, tmp_path, client
):
    status = reports.collection_status(db)
    assert status["counts"]["user_reported_located"] == 6
    assert (
        status["counts"]["codex_file_read"]
        == status["counts"]["database_imported"]
        == 0
    )
    _, _, doc, row = import_book(db, tmp_path)
    state = reports.collection_status(db)
    assert (
        state["counts"]["database_imported"] == state["counts"]["codex_file_read"] == 1
    )
    assert (
        state["counts"]["human_text_review_complete"] == 0
        and state["published_questions"] == 0
    )
    learner = {"x-api-key": "learner-test-key"}
    reviewer = {"x-review-api-key": "operator-test-key"}
    assert (
        client.get("/v2/sources/collection-status", headers=learner).status_code == 403
    )
    assert (
        client.get("/v2/sources/collection-status", headers=reviewer).status_code == 200
    )
    body = dict(review_kind="text", status="approved", reviewer="test", note=NOTE)
    result = client.post(
        f"/v2/sources/{row['id']}/reviews", headers=reviewer, json=body
    )
    assert result.status_code == 200 and result.json()["review_kind"] == "text"
    assert reports.document_report(db, doc)["stages"]["science_reviewed_units"] == 0


def test_ocr_only_published_bank_runs_28_150_60_and_pending_grading(
    db, tmp_path, question_factory, passing_audit, blueprint
):
    from app.training_spec import SUBJECTS, COUNTS, POINTS
    from app.training_sessions import start
    from app.common import ensure_user
    from app import quiz_v2 as q

    documents = {}
    styles = {}
    for subject, book, exam in [
        ("biochemistry", "biochemistry_lower", "biochemistry_exams"),
        ("cell_biology", "cell_biology", "cell_biology_exams"),
    ]:
        book_dir = tmp_path / book
        book_dir.mkdir()
        path, meta = material(book_dir, book)
        doc = sources.import_source(db, path, meta)
        unit = db.execute(
            "SELECT * FROM source_units WHERE document_id=%s ORDER BY ordinal LIMIT 1",
            (doc,),
        ).fetchone()
        approve(db, unit["id"], "text")
        approve(db, unit["id"], "science")
        documents[subject] = evidence(meta, doc, unit, sha(path.read_bytes()))
        exam_dir = tmp_path / exam
        exam_dir.mkdir()
        path, meta = material(
            exam_dir,
            exam,
            content=["2026年中国科学技术大学真题", "1. Synthetic exam style reference"],
        )
        examdoc = sources.import_source(db, path, meta)
        examunit = db.execute(
            "SELECT * FROM source_units WHERE document_id=%s AND text LIKE '1. Synthetic%%'",
            (examdoc,),
        ).fetchone()
        approve(db, examunit["id"], "text")
        approve(db, examunit["id"], "exam_style")
        scope = subject + "-ocr-style"
        exams.register_scope(
            db,
            dict(
                scope_key=scope,
                subject=subject,
                paper_code="synthetic",
                syllabus="test",
                source_unit_id=examunit["id"],
                reviewer="test",
                note=NOTE,
            ),
        )
        for kind in COUNTS:
            record = exams.add(
                db,
                dict(
                    source_unit_id=examunit["id"],
                    original_question_id=kind,
                    subject=subject,
                    year=2026,
                    question_type=kind,
                    stem="Synthetic exam style reference",
                    category="mechanism",
                ),
            )
            exams.classify(
                db,
                record["id"],
                dict(
                    scope_key=scope,
                    source_kind="transcription",
                    mapping_verified=False,
                    tag_contributions={},
                    primary_tags=[],
                    cognitive_level=None,
                    canonical_item_id=None,
                ),
                "test",
                NOTE,
            )
            styles[(subject, kind)] = record["id"]
    for subject in SUBJECTS:
        for kind, count in COUNTS.items():
            for _ in range(count):
                data = question_factory(subject, kind).model_dump()
                data["evidence"] = [documents[subject]]
                data["style_reference_ids"] = [styles[(subject, kind)]]
                question = __import__(
                    "app.question_quality", fromlist=["Question"]
                ).Question.model_validate(data)
                qid = q.create_question(db, question)["id"]
                audit = q.review_question(db, qid, passing_audit)
                assert audit["passed"], audit
                q.publish_question(db, qid)
    uid = ensure_user(db, "synthetic-ocr-only-owner")
    prepared = q.create_paper(db, uid, blueprint)
    start(db, uid, prepared["paper_id"])
    paper = q.load_paper(db, uid, prepared["paper_id"])
    assert (
        len(paper["questions"]) == 28
        and paper["max_points"] == 150
        and paper["duration_seconds"] == 3600
    )
    answers = {
        i["id"]: "A" if i["question_type"] == "single_choice" else "fixture solution"
        for i in paper["questions"]
    }
    submitted = q.submit(db, uid, paper["paper_id"], answers)
    assert submitted["status"] == "pending_review" and submitted["score"] is None
    for item in paper["questions"]:
        if item["question_type"] not in ("single_choice", "fill_blank"):
            q.grade(
                db,
                paper["paper_id"],
                item["id"],
                [POINTS[item["question_type"]]],
                NOTE,
                "test",
            )
    result = q.result(db, uid, paper["paper_id"])
    assert result["score"] == 150 and result["max_points"] == 150
    assert all(s["score"] == 75 for s in result["subject_scores"].values())
    assert result["selection_analysis"]["algorithm_version"] == "knowledge-priority/1.0"
    assert all(
        r["source_kind"] != "real_exam" or not r["trusted"]
        for r in exams.trusted_records(
            db, db.execute("SELECT clock_timestamp() AS now").fetchone()["now"]
        )
        if r["source_document_id"]
        in [
            db.execute(
                "SELECT document_id FROM source_units WHERE id=%s",
                (e["source_unit_id"],),
            ).fetchone()["document_id"]
            for e in documents.values()
        ]
    )


def test_approved_old_version_paper_score_survives_new_source_import(
    db, tmp_path, question_factory, passing_audit, blueprint, bank
):
    from app.training_sessions import start
    from app.common import ensure_user
    from app import quiz_v2 as q

    uid = ensure_user(db, "history-protection")
    paper = q.create_paper(db, uid, blueprint)
    start(db, uid, paper["paper_id"])
    q.submit(db, uid, paper["paper_id"], {})
    before = q.result(db, uid, paper["paper_id"])
    snapshots = db.execute(
        "SELECT snapshot FROM paper_items_v2 WHERE paper_id=%s ORDER BY position",
        (paper["paper_id"],),
    ).fetchall()
    path, meta = material(tmp_path, "cell_biology")
    new = sources.import_source(db, path, meta)
    assert new
    after = q.result(db, uid, paper["paper_id"])
    assert before["score"] == after["score"] == 0
    assert before["answers"] == after["answers"]
    assert (
        db.execute(
            "SELECT snapshot FROM paper_items_v2 WHERE paper_id=%s ORDER BY position",
            (paper["paper_id"],),
        ).fetchall()
        == snapshots
    )


def test_docx_table_provenance_is_retained_and_needs_review(db, tmp_path):
    path, meta = material(tmp_path, "cell_biology_exams")
    doc = sources.import_source(db, path, meta)
    row = db.execute(
        "SELECT * FROM source_units WHERE document_id=%s AND section LIKE %s",
        (doc, "%cell1%"),
    ).fetchone()
    assert row and "docx_table_requires_review" in row["quality_flags"]
    approve(db, row["id"], "text")
    with pytest.raises(ValueError, match="Unresolved"):
        approve(db, row["id"], "exam_style")


def test_collection_check_reports_only_six_local_files_without_db_or_fallback(tmp_path):
    from app.source_policy import handoff

    root = tmp_path / "OCR"
    root.mkdir()
    configured = handoff(root)
    for item in configured["files"]:
        directory = (root / item["relative_path"]).parent
        directory.mkdir(parents=True, exist_ok=True)
        material(directory, item["key"])
    report = tmp_path / "collection.json"
    assert (
        source_intake.main(
            ["collection-check", "--directory", str(root), "--report", str(report)]
        )
        == 0
    )
    value = json.loads(report.read_text())
    assert value["counts"]["file_read"] == value["counts"]["preflight_passed"] == 6
    assert (
        value["counts"]["database_imported"]
        == value["counts"]["human_text_reviewed"]
        == 0
    )
    assert sum(f["preflight"]["ocr_page_markers"] for f in value["files"]) == 1915
    PathToRemove = root / configured["files"][0]["relative_path"]
    PathToRemove.unlink()
    failed = tmp_path / "collection-missing.json"
    assert (
        source_intake.main(
            ["collection-check", "--directory", str(root), "--report", str(failed)]
        )
        == 2
    )
    assert json.loads(failed.read_text())["counts"]["file_read"] == 5


def test_pending_bibliography_can_be_reviewed_append_only_without_rewriting_source(
    db, tmp_path, question_factory
):
    path, meta = material(tmp_path)
    meta["authors"] = None
    meta["edition"] = "待核查"
    doc = sources.import_source(db, path, meta)
    unit = db.execute(
        "SELECT * FROM source_units WHERE document_id=%s ORDER BY ordinal LIMIT 1",
        (doc,),
    ).fetchone()
    approve(db, unit["id"], "text")
    approve(db, unit["id"], "science")
    before = db.execute("SELECT * FROM source_documents WHERE id=%s", (doc,)).fetchone()
    book = dict(
        title="Actual reviewed synthetic title",
        authors="Test author",
        edition="Test edition",
        isbn=None,
    )
    event = ocr.review_bibliography(db, doc, book, "test", NOTE)
    assert (
        db.execute("SELECT * FROM source_documents WHERE id=%s", (doc,)).fetchone()
        == before
    )
    question = question_factory().model_dump()
    actual = {
        **meta,
        "title": book["title"],
        "authors": book["authors"],
        "edition": book["edition"],
    }
    question["evidence"] = [evidence(actual, doc, unit, sha(path.read_bytes()))]
    assert not sources.evidence_issues(db, question)
    with pytest.raises(RaiseException), db.transaction():
        db.execute(
            "DELETE FROM source_bibliography_reviews WHERE id=%s", (event["id"],)
        )


def test_old_hash_cannot_silently_reactivate_its_approved_evidence(db, tmp_path):
    path, meta, doc, row = import_book(db, tmp_path)
    original = path.read_bytes()
    approve(db, row["id"], "text")
    path.write_text(path.read_text().replace("paragraph 1", "updated paragraph 1", 1))
    sources.import_source(db, path, meta)
    path.write_bytes(original)
    with pytest.raises(ValueError, match="reactivate"):
        sources.import_source(db, path, meta)


@pytest.mark.parametrize("key", ["biochemistry_upper", "biochemistry_lower"])
def test_labelled_backslash_markers_also_supported(tmp_path, key):
    path, meta = material(tmp_path, key)
    path.write_text(path.read_text().replace("===== ", r"\===== ", 1))
    result = sources.prepare_source(path, meta)
    assert result["pages"][0]["marker_text"].startswith("\\===== ")


def test_ocr_transcripts_cannot_be_reclassified_as_original_exam(db, tmp_path):
    path, meta = material(tmp_path, "cell_biology_exams")
    doc = sources.import_source(db, path, meta)
    unit = db.execute(
        "SELECT * FROM source_units WHERE document_id=%s AND text LIKE '1. Synthetic%%' ORDER BY id LIMIT 1",
        (doc,),
    ).fetchone()
    approve(db, unit["id"], "text")
    scope = "ocr-scope"
    exams.register_scope(
        db,
        dict(
            scope_key=scope,
            subject="cell_biology",
            paper_code="synthetic",
            syllabus="test",
            source_unit_id=unit["id"],
            reviewer="test",
            note=NOTE,
        ),
    )
    record = exams.add(
        db,
        dict(
            source_unit_id=unit["id"],
            original_question_id="test",
            subject="cell_biology",
            year=2023,
            question_type="single_choice",
            stem=unit["text"],
            category="mechanism",
        ),
    )
    with pytest.raises(ValueError, match="transcription"):
        exams.classify(
            db,
            record["id"],
            dict(
                scope_key=scope,
                source_kind="real_exam",
                mapping_verified=False,
                tag_contributions={},
                primary_tags=[],
                cognitive_level=None,
                canonical_item_id=None,
            ),
            "test",
            NOTE,
        )


def test_current_annual_report_excludes_outside_collection_records(db, tmp_path):
    path = tmp_path / "old-source.txt"
    path.write_text("[PAGE 1]\nSynthetic old question")
    doc = sources.import_source(
        db,
        path,
        dict(
            file_id="old-superseded-id",
            title="Old",
            edition="Test",
            subject="biochemistry",
            expected_pages=1,
        ),
    )
    unit = db.execute(
        "SELECT id FROM source_units WHERE document_id=%s", (doc,)
    ).fetchone()["id"]
    exams.add(
        db,
        dict(
            source_unit_id=unit,
            original_question_id="old",
            subject="biochemistry",
            year=2026,
            question_type="essay",
            stem="Synthetic old question",
            category="mechanism",
        ),
    )
    row = next(
        r
        for r in reports.annual_report(db)
        if r["subject"] == "biochemistry" and r["year"] == 2026
    )
    assert (
        row["questions"] == row["candidate_questions"] == 0
        and row["status"] == "original_not_acquired"
    )


def test_pdf_tool_cannot_impersonate_a_current_txt_identity(tmp_path):
    path, meta = material(tmp_path)
    meta.pop("source_mode")
    with pytest.raises(ValueError, match="cannot impersonate"):
        sources.prepare_source(path, meta)


def test_reviewed_docx_record_does_not_claim_original_exam_acquired(db, tmp_path):
    path, meta = material(tmp_path, "biochemistry_exams")
    doc = sources.import_source(db, path, meta)
    unit = db.execute(
        "SELECT * FROM source_units WHERE document_id=%s AND text LIKE '1. Synthetic%%' ORDER BY id LIMIT 1",
        (doc,),
    ).fetchone()
    approve(db, unit["id"], "text")
    exams.add(
        db,
        dict(
            source_unit_id=unit["id"],
            original_question_id="transcript",
            subject="biochemistry",
            year=2023,
            question_type="single_choice",
            stem=unit["text"],
            category="mechanism",
        ),
    )
    row = next(
        r
        for r in reports.annual_report(db)
        if r["subject"] == "biochemistry" and r["year"] == 2023
    )
    assert row["ocr_transcription_file_acquired"]
    assert not row["original_acquired"] and not row["original_verified"]
    assert not row["included_in_reliable_frequency"]


def test_unknown_volume_label_is_an_invalid_marker_not_body_text():
    with pytest.raises(ValueError, match="Malformed"):
        sources.parse("===== 中册 page-001 =====\nUnlocated synthetic text", {})


def test_embedded_docx_image_without_text_is_indexed_as_unresolved_not_invented(
    db, tmp_path
):
    import io
    from PIL import Image

    path, meta = material(tmp_path, "cell_biology_exams")
    doc = Document(path)
    picture = io.BytesIO()
    Image.new("RGB", (4, 4), "white").save(picture, format="PNG")
    picture.seek(0)
    doc.add_picture(picture)
    doc.save(path)
    document = sources.import_source(db, path, meta)
    row = db.execute(
        "SELECT * FROM source_units WHERE document_id=%s AND quality_flags ? 'untranscribed_image'",
        (document,),
    ).fetchone()
    assert row and row["text"] == "[DOCX image without transcribed text]"
    with pytest.raises(ValueError, match="Untranscribed"):
        approve(db, row["id"], "text")


def test_future_source_revision_does_not_rewrite_historical_asof_inputs(
    db, tmp_path, question_factory
):
    cutoff = db.execute("SELECT clock_timestamp() AS now").fetchone()["now"]
    before = {
        r["id"]
        for r in exams.trusted_records(db, cutoff)
        if r["trusted"] and r["subject"] == "cell_biology"
    }
    assert before  # Synthetic pre-existing scoped source, not a real exam claim.
    path, meta = material(tmp_path, "cell_biology")
    sources.import_source(db, path, meta)
    current = db.execute("SELECT clock_timestamp() AS now").fetchone()["now"]
    assert not any(
        r["trusted"] for r in exams.trusted_records(db, current) if r["id"] in before
    )
    past = {
        r["id"]
        for r in exams.trusted_records(db, cutoff)
        if r["trusted"] and r["subject"] == "cell_biology"
    }
    assert past == before


def test_malformed_docx_is_reported_as_rejected_not_imported(tmp_path):
    path, meta = material(tmp_path, "cell_biology_exams")
    path.write_bytes(b"not a DOCX container")
    metadata = tmp_path / "metadata.json"
    metadata.write_text(json.dumps(meta))
    report = tmp_path / "rejected-docx.json"
    assert (
        source_intake.main(
            ["check", str(path), "--metadata", str(metadata), "--report", str(report)]
        )
        == 2
    )
    assert json.loads(report.read_text())["status"] == "rejected"
