"""Synthetic regression fixtures matching private handoff structure, never private bodies."""

import json
import stat
import zipfile

import pytest
from psycopg.errors import RaiseException
from app.handoff_bundle import unpack
from app.source_files import sha
from app.source_policy import COLLECTION, FILES
from app.source_paging import PAGE
from app.source_structure import book_layout
from app.transcript_analysis import candidates, year_heading
from app.pending_mappings import persist
from test_ocr_only import material
from app import sources
from app.local_corpus import mapping_proposals, year_report
from app.ocr_sources import docx_blocks
from docx import Document


def bundle(path, *, corrupt=False, extra=None):
    entries = []
    with zipfile.ZipFile(path, "w") as z:
        for n, item in enumerate(FILES.values()):
            raw = f"SYNTHETIC PRIVATE HANDOFF {n}".encode()
            member = "files/" + item["relative_path"]
            z.writestr(member, raw)
            entries.append(
                dict(
                    google_drive_file_id=item["file_id"],
                    filename=item["filename"],
                    archive_path=member,
                    bytes=len(raw),
                    sha256="0" * 64 if corrupt else sha(raw),
                )
            )
        z.writestr(
            "manifest.json",
            json.dumps(
                dict(
                    schema="private-ocr-handoff/v1",
                    source_folder_id=COLLECTION["folder_id"],
                    files=entries,
                )
            ),
        )
        if extra:
            z.writestr(extra, b"extra")


def test_bundle_all_hashes_before_extraction_and_private_permissions(tmp_path):
    archive = tmp_path / "handoff.zip"
    bundle(archive)
    out = tmp_path / "private"
    report = unpack(archive, out)
    assert report["hash_matches"] == 6 and report["database_imported"] == 0
    assert stat.S_IMODE(out.stat().st_mode) == 0o700
    for file in out.rglob("*"):
        if file.is_file():
            assert stat.S_IMODE(file.stat().st_mode) == 0o600
    with pytest.raises(ValueError, match="new"):
        unpack(archive, out)


@pytest.mark.parametrize(
    "extra", ["../escape.txt", "/absolute.txt", "files/foreign.txt"]
)
def test_bundle_rejects_unsafe_or_foreign_files_before_writing(tmp_path, extra):
    archive = tmp_path / "bad.zip"
    bundle(archive, extra=extra)
    with pytest.raises(ValueError):
        unpack(archive, tmp_path / "out")
    assert not (tmp_path / "out").exists()


def test_bundle_checksum_failure_extracts_nothing(tmp_path):
    archive = tmp_path / "bad.zip"
    bundle(archive, corrupt=True)
    with pytest.raises(ValueError, match="checksum"):
        unpack(archive, tmp_path / "out")
    assert not (tmp_path / "out").exists()


def test_numeric_toc_aligns_body_without_using_printed_page_as_ocr_page():
    text = "\n".join(
        [
            "===== 上册 page-001 =====",
            "1 合成章节 9",
            "1.1 合成小节标题 10",
            "2 另一个章节 20",
            "===== 上册 page-002 =====",
            "1",
            "合成章节",
            "1.1 合成小节标题",
            "Synthetic body",
        ]
    )
    layout = book_layout(text, PAGE)
    assert layout["toc_pages"] == {1}
    assert layout["headings"][5] == "1 合成章节"
    assert layout["headings"][7] == "1 合成章节"
    assert layout["catalogue"][0]["toc_ocr_page"] == 1


def parsed(texts, sections=None):
    units = [
        dict(ordinal=n, text=t, section=(sections or {}).get(n, ""))
        for n, t in enumerate(texts, 1)
    ]
    return list(candidates(units, dict(year_start=2004, year_end=2026)))


def test_native_word_numbering_options_subquestions_and_answer_boundaries():
    rows = parsed(
        [
            "2026年合成考试真题",
            "一、名词解释",
            "5.8SrRNA",
            "Synthetic option",
            "Synthetic nested",
            "参考答案：仅测试",
            "1. Answer",
        ],
        {
            3: "p1|num:1:0:decimal:1",
            4: "p2|num:1:1:lowerLetter:1",
            5: "p3|num:1:1:decimal:1",
        },
    )
    assert [r["analysis"]["role"] for r in rows] == [
        "year_heading",
        "section_heading",
        "question_start",
        "option",
        "subquestion",
        "answer",
        "answer",
    ]
    assert rows[2]["question_number"] == "1"
    assert rows[-1]["analysis"]["parent_question_unit_id"] is None


def test_scientific_term_without_native_number_is_not_question_five():
    rows = parsed(["2026年合成考试真题", "名词解释", "5.8SrRNA", "1. 解析某机制"])
    assert rows[2]["question_number"] is None
    assert rows[3]["analysis"]["role"] == "question_start"


def test_private_year_report_separates_answers_and_never_claims_original_verified():
    file = dict(
        file_id="synthetic",
        sha256="a" * 64,
        year_start=2026,
        year_end=2026,
        rows=[
            dict(ordinal=n, text=text, section="", ocr_global_page=None)
            for n, text in enumerate(
                [
                    "2026年合成考试真题",
                    "名词解释",
                    "1. Synthetic question",
                    "参考答案：Synthetic answer",
                    "2. Another answer",
                ],
                1,
            )
        ],
    )
    report, questions = year_report(file)
    assert len(questions) == 1 and questions[0]["text"] == "1. Synthetic question"
    year = report["years"][0]
    assert year["candidate_questions"] == 1
    assert (
        year["reference_answer_markers"] == 2
    )  # answer paragraphs, not proof of official answers
    assert not year["original_exam_verified"] and not year["trusted_frequency_sample"]


def test_literal_mapping_does_not_match_atp_inside_atpase_or_approve_terms():
    book = dict(
        file_id="book",
        sha256="b" * 64,
        kind="textbook",
        subject="biochemistry",
        rows=[
            dict(
                ordinal=1,
                text="合成术语(ATP)",
                chapter="合成章",
                section="",
                ocr_global_page=1,
                quality_flags=[],
            )
        ],
    )
    exam = dict(file_id="exam", kind="exam_transcription", subject="biochemistry")
    question = dict(
        file_id="exam",
        source_sha256="c" * 64,
        ordinal=1,
        ocr_page=None,
        local_unit_id="c:1",
        text="ATPase",
    )
    proposals = mapping_proposals([book, exam], [question])
    assert all(p["target"] is None and p["status"] == "pending" for p in proposals)
    proposals = mapping_proposals([book, exam], [{**question, "text": "ATP"}])
    assert any(p["target"] is not None for p in proposals)
    assert all(p["status"] == "pending" for p in proposals)


@pytest.mark.parametrize(
    "text", ["2004—2026年真题汇编", "解释2026年实验", "2026年真题为何如此？"]
)
def test_year_mentions_do_not_establish_year_coverage(text):
    assert year_heading(text, dict(year_start=2004, year_end=2026)) is None


def test_expected_manifest_hash_is_checked_before_source_import(tmp_path):
    path, metadata = material(tmp_path)
    metadata["expected_sha256"] = "0" * 64
    with pytest.raises(ValueError, match="SHA|hash|checksum"):
        sources.prepare_source(path, metadata)


def test_docx_style_numbering_is_recovered_without_changing_original_text(tmp_path):
    doc = Document()
    doc.add_paragraph("2026年合成考试真题")
    doc.add_paragraph("名词解释")
    doc.add_paragraph("5.8SrRNA", style="List Number")
    doc.add_paragraph("Synthetic second\nembedded continuation", style="List Number")
    path = tmp_path / "numbered.docx"
    doc.save(path)
    blocks = docx_blocks(path.read_bytes())
    assert blocks[2][0] == "5.8SrRNA"
    assert blocks[2][-1]["value"] == 1
    assert blocks[3][-1]["value"] == 2
    assert blocks[2][-1]["format"] == "decimal"


def test_pending_mappings_are_idempotent_immutable_and_not_authoritative(db, tmp_path):
    path, meta = material(tmp_path, content="Synthetic kinase evidence")
    sources.import_source(db, path, meta)
    ref = dict(
        file_id=meta["file_id"],
        source_sha256=sha(path.read_bytes()),
        ordinal=1,
        ocr_page=1,
    )
    proposal = dict(
        source=ref,
        target=None,
        knowledge_label="kinase",
        method="test",
        status="pending",
        evidence={"quote": "Synthetic kinase evidence"},
    )
    assert persist(db, [proposal]) == 1
    assert persist(db, [proposal]) == 0
    assert db.execute("SELECT count(*) AS n FROM knowledge_links").fetchone()["n"] == 0
    with pytest.raises(RaiseException), db.transaction():
        db.execute("UPDATE source_mapping_proposals SET method='tampered'")
    with pytest.raises(ValueError, match="quote"):
        persist(db, [{**proposal, "evidence": {"quote": "invented"}}])
    with pytest.raises(ValueError, match="identity"):
        persist(db, [{**proposal, "source": {**ref, "source_sha256": "0" * 64}}])
