"""Automatically select authentic user-provided historical exam wording.

Synthetic test DOCX fixtures prove a previous manually seeded review packet is
not required, but must never be mistaken for a production past-exam source.
"""
from app import sources
from app.auto_generation import select_exam_style
from app.candidate_review import snapshot_unit
from test_ocr_only import material


def test_style_anchor_reads_real_exam_transcript_without_candidate_packet(db, tmp_path):
    path, metadata = material(tmp_path, "biochemistry_exams", content=[
        "2023年中国科学技术大学生物化学真题",
        "单项选择题",
        "1. Synthetic 2023 question",
        "A. Option one B. Option two C. Option three D. Option four",
        "2024年中国科学技术大学生物化学真题",
        "简答题",
        "1. Synthetic 2024 question",
        "2025年中国科学技术大学生物化学真题",
    ])
    sources.import_source(db, path, metadata)
    assert db.execute("SELECT count(*) AS n FROM candidate_review_packets").fetchone()["n"] == 0
    style = select_exam_style(db, "biochemistry", "single_choice")
    assert style and style["year"] == 2023
    paragraph = snapshot_unit(db, style["id"])
    assert paragraph["source_kind"] == "exam_transcription"
    assert "Synthetic 2023 question" in paragraph["text"]
    assert select_exam_style(db, "biochemistry", "short_answer")["year"] == 2024
    assert select_exam_style(db, "biochemistry", "fill_blank") is None
    assert select_exam_style(db, "cell_biology", "single_choice") is None


def test_explicit_docx_fidelity_does_not_invent_missing_exam_question_types(db, tmp_path):
    path, metadata = material(tmp_path, "cell_biology_exams", content=[
        "2025年中国科学技术大学细胞生物学真题",
        "名词解释",
        "1. Synthetic cell-biology definition",
        "2026年中国科学技术大学细胞生物学真题",
    ])
    sources.import_source(db, path, metadata)
    result = select_exam_style(db, "cell_biology", "term")
    assert result and result["year"] == 2025
    assert select_exam_style(db, "cell_biology", "single_choice") is None


def test_style_example_prefers_same_actual_topic_without_inferring_scientific_answer(db, tmp_path):
    path, metadata = material(tmp_path, "biochemistry_exams", content=[
        "2023年中国科学技术大学生物化学真题",
        "单项选择题",
        "1. Synthetic ATP-dependent enzyme example",
        "A. First B. Second C. Third D. Fourth",
        "2024年中国科学技术大学生物化学真题",
        "单项选择题",
        "1. Synthetic SRP-associated targeting example",
        "A. Alpha B. Beta C. Gamma D. Delta",
        "2025年中国科学技术大学生物化学真题",
    ])
    sources.import_source(db, path, metadata)
    latest = select_exam_style(db, "biochemistry", "single_choice")
    assert latest["year"] == 2024
    matched = select_exam_style(db, "biochemistry", "single_choice", {
        "unit_id": 7, "relation": {"name": "ATP", "target": ""}
    })
    assert matched["year"] == 2023
    assert matched["lexical_topic_overlap"] == 1
    assert matched["scope"] == "past_exam_style_not_textbook_answer_evidence"
    assert "answer" not in matched


def test_separate_untranscribed_figure_is_excluded_from_generation_style(db, tmp_path):
    from docx import Document
    from PIL import Image

    path, metadata = material(tmp_path, "biochemistry_exams", content=[
        "2025年中国科学技术大学生物化学真题",
        "单项选择题",
        "1. Examine the synthetic diagram below",
        "2. Complete self-contained synthetic text question",
        "A. First B. Second C. Third D. Fourth",
        "2026年中国科学技术大学生物化学真题",
    ])
    image = tmp_path / "synthetic-exam-figure.png"
    Image.new("RGB", (10, 10), "white").save(image)
    document = Document(path)
    # Insert the image immediately after Q1, not within its own question stem.
    # Move the newly appended drawing paragraph immediately after Q1.
    # The original question text does not itself carry the image flag.
    document.add_picture(str(image))
    image_paragraph = document.paragraphs[-1]
    document.paragraphs[2]._p.addnext(image_paragraph._p)
    document.save(path)

    sources.import_source(db, path, metadata)
    q1 = db.execute(
        "SELECT concerns FROM exam_transcript_candidates "
        "WHERE transcription LIKE '1. Examine%'"
    ).fetchone()
    assert q1 is not None and "untranscribed_image" in q1["concerns"]
    chosen = select_exam_style(db, "biochemistry", "single_choice")
    assert chosen is not None
    assert "Complete self-contained" in chosen["transcription"]
    assert "Examine" not in chosen["transcription"]


def test_critical_following_docx_table_blocks_parent_style_question(db, tmp_path):
    """A question in normal text cannot silently lose its following table."""
    from docx import Document

    path, metadata = material(tmp_path, "biochemistry_exams", content=[
        "2025年中国科学技术大学生物化学真题",
        "单项选择题",
        "1. Interpret the synthetic results in the following table",
        "2. Self-contained synthetic choice question",
        "A. First B. Second C. Third D. Fourth",
        "2026年中国科学技术大学生物化学真题",
    ])
    doc = Document(path)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Critical tabulated synthetic results"
    # Move table between Q1 and Q2 (not an image and not embedded in Q1).
    doc.paragraphs[2]._p.addnext(table._tbl)
    doc.save(path)

    sources.import_source(db, path, metadata)
    q1 = db.execute(
        "SELECT concerns FROM exam_transcript_candidates "
        "WHERE transcription LIKE '1. Interpret%'"
    ).fetchone()
    assert q1 is not None
    assert "docx_table_requires_review" in q1["concerns"]
    candidate = select_exam_style(db, "biochemistry", "single_choice")
    assert candidate is not None
    assert "Self-contained" in candidate["transcription"]
    assert "Interpret" not in candidate["transcription"]


def test_generated_v2_prompt_explicitly_binds_scoring_and_evidence():
    from app.exam_authoring import instructions
    for kind, points in (("single_choice", 2), ("fill_blank", 2),
                         ("term", 4), ("short_answer", 9)):
        prompt = instructions(kind, points)
        assert f"EXACTLY {points} points" in prompt
        assert "verbatim continuous substring" in prompt
        assert "source_sha256" in prompt
        assert "JSON integer" in prompt
        assert "sum is exactly max_points" in prompt
        assert "single unwrapped JSON" in prompt
