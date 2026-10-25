import pytest
from app import sources, exams, quiz_v2


def test_385_page_volume_mapping_is_explicit():
    text = "\n".join(f"[PAGE {i}]\nVolume one paragraph {i}\n" for i in range(1, 193))
    text += "\n".join(f"[PAGE {i}]\nVolume two paragraph {i}\n" for i in range(1, 194))
    rows = sources.parse(
        text,
        {"volume_offsets": [0, 192], "volume_pages": [192, 193], "expected_pages": 385},
    )
    assert (
        len(rows) == 385 and rows[192]["ocr_page"] == 1 and rows[192]["pdf_page"] == 193
    )
    assert rows[-1]["pdf_page"] == 385 and rows[-1]["volume"] == 2
    with pytest.raises(ValueError, match="reset"):
        sources.parse(text, {})
    with pytest.raises(ValueError, match="coverage"):
        sources.parse("[PAGE 1]\ntext", {"expected_pages": 2})


def test_private_import_idempotence_and_unverified_default(db, tmp_path):
    text = tmp_path / "book.txt"
    text.write_text("[PAGE 1]\n第一章 酶学\nKm and Vmax are scientific notation.\n")
    metadata = dict(
        file_id="synthetic-id",
        title="Synthetic",
        edition="Test",
        subject="biochemistry",
        expected_pages=1,
    )
    doc = sources.import_source(db, text, metadata)
    assert sources.import_source(db, text, metadata) == doc
    unit = sources.search(db, "Km")[0]
    assert unit["quality_flags"] and not unit["verified"] and unit["pdf_page"] == 1
    with pytest.raises(ValueError, match="missing"):
        sources.verify_unit(
            db, unit["id"], "reviewer", "Detailed original verification note"
        )
    assert db.execute("SELECT COUNT(*) AS n FROM source_documents").fetchone()["n"] == 1


def test_corrupt_ocr_not_verified(db, tmp_path):
    raw = tmp_path / "bad.txt"
    raw.write_text("[PAGE 1]\nCorrupt \ufffd OCR")
    original = tmp_path / "original.bin"
    original.write_bytes(b"TEST ORIGINAL")
    doc = sources.import_source(
        db,
        raw,
        dict(file_id="bad", title="Bad", edition="Test", subject="cell_biology"),
        original,
    )
    unit = db.execute(
        "SELECT id FROM source_units WHERE document_id=%s", (doc,)
    ).fetchone()["id"]
    with pytest.raises(ValueError, match="Corrupt"):
        sources.verify_unit(
            db, unit, "reviewer", "I compared this paragraph with the original"
        )


def test_source_evidence_page_text_and_original_checks(
    db, question_factory, passing_audit
):
    q = question_factory()
    unit = q.evidence[0].source_unit_id
    db.execute("UPDATE source_units SET pdf_page=3 WHERE id=%s", (unit,))
    assert "textbook_page_mismatch" in sources.evidence_issues(db, q.model_dump())
    q.evidence[0].locator = "PDF:3"
    assert not sources.evidence_issues(db, q.model_dump())
    q.evidence[0].excerpt = "Invented quotation"
    assert "textbook_evidence_mismatch" in sources.evidence_issues(db, q.model_dump())
    qid = quiz_v2.create_question(db, q)["id"]
    assert not quiz_v2.review_question(db, qid, passing_audit)["passed"]
    with pytest.raises(ValueError):
        quiz_v2.publish_question(db, qid)
    db.execute("UPDATE source_documents SET original_content=NULL")
    assert "missing_or_changed_original" in sources.evidence_issues(db, q.model_dump())


def test_quality_score_and_style_evidence_are_hard_gates(
    db, question_factory, passing_audit
):
    q = question_factory()
    q.style_reference_ids = []
    qid = quiz_v2.create_question(db, q)["id"]
    audit = passing_audit.model_copy(update={"score": 89})
    result = quiz_v2.review_question(db, qid, audit)
    assert (
        "missing_exam_style_source" in result["issues"]
        and "quality_score_below_90" in result["issues"]
    )


def test_exam_year_coverage_and_duplicate_sources(db, question_factory):
    q = question_factory()
    unit = q.evidence[0].source_unit_id
    db.execute(
        "UPDATE source_units SET text=text || ' Synthetic original exam question' WHERE id=%s",
        (unit,),
    )
    record = dict(
        source_unit_id=unit,
        original_question_id="source-1",
        subject="biochemistry",
        year=2024,
        question_type="single_choice",
        stem="Synthetic original exam question",
        category="mechanism",
        knowledge_tags=["Km"],
        original_verified=True,
    )
    assert exams.add(db, record)
    assert exams.add(db, record) is None
    record["original_question_id"] = "different-transcription"
    exams.add(db, record)
    rates = exams.frequencies(db)
    assert next(r for r in rates if r["tag"] == "Km")["appearances"] == 1
    coverage = exams.coverage(db)
    missing = next(
        r for r in coverage if r["subject"] == "biochemistry" and r["year"] == 2026
    )
    assert missing["status"] == "original_not_acquired"
    covered = next(
        r for r in coverage if r["subject"] == "biochemistry" and r["year"] == 2024
    )
    assert (
        covered["missing_answers"]
        and covered["missing_textbook_evidence"]
        and not covered["complete_original"]
    )
    record["answer_kind"] = "official"
    record["answer"] = "A"
    with pytest.raises(ValueError, match="separate"):
        exams.add(db, record)


def test_txt_docx_pdf_import_extractors(tmp_path):
    from docx import Document
    from pypdf import PdfWriter

    doc = Document()
    doc.add_paragraph("Synthetic DOCX text")
    doc.add_table(rows=1, cols=1).cell(
        0, 0
    ).text = "Synthetic exam question inside table"
    doc.save(tmp_path / "source.docx")
    assert "Synthetic exam question inside table" in sources.extract(
        (tmp_path / "source.docx").read_bytes(), ".docx"
    )
    assert "Synthetic DOCX text" in sources.extract(
        (tmp_path / "source.docx").read_bytes(), ".docx"
    )
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    writer.write(tmp_path / "source.pdf")
    with pytest.raises(ValueError, match="No text"):
        sources.parse(
            sources.extract((tmp_path / "source.pdf").read_bytes(), ".pdf"), {}
        )
    with pytest.raises(ValueError):
        sources.extract(b"anything", ".exe")


def test_bibliography_cannot_reference_another_isbn_or_author(db, question_factory):
    q = question_factory()
    q.evidence[0].authors = "Invented author"
    assert "textbook_bibliography_mismatch" in sources.evidence_issues(
        db, q.model_dump()
    )


def test_later_original_attachment_preserves_unit_identity_and_checks_physical_pages(
    db, tmp_path
):
    from pypdf import PdfWriter

    text = tmp_path / "ocr.txt"
    text.write_text("[PAGE 1]\nSynthetic paragraph")
    metadata = dict(
        file_id="attach-later",
        title="Test",
        edition="Test",
        subject="biochemistry",
        pdf_pages={"1": 2},
    )
    doc = sources.import_source(db, text, metadata)
    unit = db.execute(
        "SELECT id FROM source_units WHERE document_id=%s", (doc,)
    ).fetchone()["id"]
    original = tmp_path / "original.pdf"
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    writer.add_blank_page(100, 100)
    writer.write(original)
    assert sources.import_source(db, text, metadata, original) == doc
    assert (
        db.execute(
            "SELECT id FROM source_units WHERE document_id=%s", (doc,)
        ).fetchone()["id"]
        == unit
    )
    sources.verify_unit(
        db,
        unit,
        "test-reviewer",
        "Synthetic original page mapping reviewed for this test",
    )
    with pytest.raises(ValueError, match="immutable"):
        sources.import_source(
            db, text, {**metadata, "title": "Changed title"}, original
        )
    with pytest.raises(ValueError, match="outside"):
        sources.import_source(
            db,
            text,
            {**metadata, "file_id": "invalid-page", "pdf_pages": {"1": 3}},
            original,
        )
